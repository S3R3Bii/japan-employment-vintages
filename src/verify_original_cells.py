"""Separate Openpyxl/calendar reconciliation of the preserved native-cell projection."""

import argparse, csv, hashlib, json, zipfile
from pathlib import Path
from decimal import Decimal
from xml.etree import ElementTree as E
import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--private-dir", type=Path, required=True)
    a = p.parse_args()
    projection = list(csv.DictReader((ROOT / "data/counts.csv").open(encoding="utf-8")))
    receipts = json.loads(
        (ROOT / "data/source_receipts.json").read_text(encoding="utf-8")
    )
    checks = 0
    before = {
        r["file"]: hashlib.sha256((a.private_dir / r["file"]).read_bytes()).hexdigest()
        for r in receipts
    }
    for receipt in receipts:
        path = a.private_dir / receipt["file"]
        assert before[receipt["file"]] == receipt["sha256"]
        checks += 1
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        chosen = [r for r in projection if r["source_file"] == receipt["file"]]
        assert len(chosen) == 144
        checks += 1
        ws = wb[chosen[0]["sheet"]]
        allrows = list(ws.iter_rows(values_only=True))
        # Every year is anchored by its explicit Gregorian February label, independent of era parsing.
        anchors = {}
        for ix, row in enumerate(allrows, 1):
            if row[2] == "Feb." and isinstance(row[1], str):
                pieces = row[1].split()
                if len(pieces) == 2 and pieces[0].isdecimal() and len(pieces[0]) == 4:
                    anchors[int(pieces[0])] = ix
        for row in chosen:
            year, month = map(int, row["period"].split("-"))
            ridx = anchors[year] + month - 2
            assert ridx == int("".join(c for c in row["cell"] if c.isdigit()))
            checks += 1
            col = 6 if row["region"] == "Southern-Kanto" else 12
            val = allrows[ridx - 1][col - 1]
            assert type(val) is int and Decimal(val) == Decimal(row["value_token"])
            checks += 1
            assert (
                row["month_label"] == allrows[ridx - 1][2]
                and row["period_label"] == allrows[ridx - 1][1]
            )
            checks += 2
            assert (
                allrows[9][col - 1] == "Employed person"
                and allrows[7][col - 1] == "就業者"
            )
            checks += 2
            assert row["cell"] == f"{openpyxl.utils.get_column_letter(col)}{ridx}"
            checks += 1
        with zipfile.ZipFile(path) as z:
            ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            native = {}
            for name in z.namelist():
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                    tree = E.fromstring(z.read(name))
                    d = {
                        c.attrib["r"]: c
                        for c in tree.findall(".//x:sheetData/x:row/x:c", ns)
                    }
                    if "F10" in d:
                        # Identify the bound sheet by the unique full selected-value vector, then require one match.
                        if all(
                            r["cell"] in d
                            and d[r["cell"]].find("x:v", ns) is not None
                            and d[r["cell"]].find("x:v", ns).text == r["value_token"]
                            for r in chosen
                        ):
                            native[name] = d
            assert len(native) == 1
            checks += 1
            d = next(iter(native.values()))
            for row in chosen:
                c = d[row["cell"]]
                assert c.attrib.get("t", "n") == row["native_type"] == "n"
                assert c.find("x:f", ns) is None
                checks += 2
        wb.close()
    after = {
        r["file"]: hashlib.sha256((a.private_dir / r["file"]).read_bytes()).hexdigest()
        for r in receipts
    }
    assert before == after
    checks += 2
    result = dict(
        status="PASS_INDEPENDENT_ORIGINAL_RECONCILIATION",
        count_cells=288,
        checks=checks,
        readers=["Openpyxl read-only", "stdlib raw OOXML tokens"],
        original_hashes_unchanged=True,
    )
    (ROOT / "data/original_verification.json").write_bytes(
        (json.dumps(result, indent=2) + "\n").encode()
    )
    print(json.dumps(result))


if __name__ == "__main__":
    main()
